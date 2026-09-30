"""
Multimodal Input Pipeline for Nexora AI.

Architectural Flow:
  Input
  -> Classify Input (text_pdf, scanned_pdf, document_image, visual_image, structured_doc)
  -> Choose Extraction Strategy:
       * Text PDFs / Docs: Native digital text extraction
       * Scanned PDFs: Page rasterization -> OCR + vision transcription
       * Document Images (receipts, assignments, forms): OCR + Vision transcription
       * Normal Photographs / Diagrams / Human Images: Direct Multimodal Vision Understanding
  -> Produce Normalized Multimodal Context (MultimodalPage / MultimodalContent)
  -> Return Context to DocumentService / Orchestrator.

Crucial Rule:
  "No OCR text extracted" does NOT mean "no useful image information."
  Image Understanding and Text Extraction are strictly separated.
"""
import os
from enum import Enum
from pathlib import Path
from typing import List, Optional, Tuple

from PIL import Image

from backend.config import settings
from backend.logger import logger
from backend.services.ocr_vision_service import (
    ocr_vision_service,
    ocr_image,
    render_pdf_page_to_image,
    is_text_sufficient,
)

IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp"}


class InputType(str, Enum):
    TEXT_PDF = "text_pdf"
    SCANNED_PDF = "scanned_pdf"
    DOCUMENT_IMAGE = "document_image"
    VISUAL_IMAGE = "visual_image"
    DIAGRAM_CHART = "diagram_chart"
    STRUCTURED_DOC = "structured_doc"


class MultimodalPage:
    def __init__(
        self,
        page_number: int,
        text: str,
        source_type: str,
        input_type: str,
        has_text: bool = True,
        has_visuals: bool = False,
        visual_description: Optional[str] = None,
    ):
        self.page_number = page_number
        self.text = text
        self.source_type = source_type  # 'text' | 'ocr' | 'vision' | 'visual_understanding' | 'table' | 'visual_media'
        self.input_type = input_type
        self.has_text = has_text
        self.has_visuals = has_visuals
        self.visual_description = visual_description


class MultimodalService:
    def classify_image(self, image: Image.Image) -> Tuple[InputType, str]:
        """Classify an image as document_image (contains text) or visual_image (photograph/scene)."""
        ocr_sample = ocr_image(image)
        if is_text_sufficient(ocr_sample):
            return InputType.DOCUMENT_IMAGE, ocr_sample
        return InputType.VISUAL_IMAGE, ocr_sample

    def process_pdf(self, path: str) -> List[MultimodalPage]:
        """Process PDF: use digital text for native pages; rasterize + OCR/vision for scans."""
        import fitz  # PyMuPDF

        pages: List[MultimodalPage] = []
        doc = fitz.open(path)
        try:
            for i in range(len(doc)):
                raw_text = doc[i].get_text().strip()
                if is_text_sufficient(raw_text):
                    pages.append(
                        MultimodalPage(
                            page_number=i + 1,
                            text=raw_text,
                            source_type="text",
                            input_type=InputType.TEXT_PDF.value,
                            has_text=True,
                            has_visuals=False,
                        )
                    )
                else:
                    # Scanned page or rasterized content: render and transcribe
                    text, source_type = ocr_vision_service.extract_page_text(path, i, raw_text)
                    pages.append(
                        MultimodalPage(
                            page_number=i + 1,
                            text=text,
                            source_type=source_type,
                            input_type=InputType.SCANNED_PDF.value,
                            has_text=bool(text.strip()),
                            has_visuals=True,
                        )
                    )
        finally:
            doc.close()
        return pages

    def process_image(self, path: str) -> List[MultimodalPage]:
        """Process standalone image (jpg/png/webp):
        - Document image: extract structured text using OCR + Vision.
        - Normal photograph / visual image: send directly to vision model for Image Understanding.
        """
        image = Image.open(path).convert("RGB")
        input_type, ocr_text = self.classify_image(image)

        if input_type == InputType.DOCUMENT_IMAGE:
            # Document-like image: text extraction + vision transcription
            text, source_type = ocr_vision_service.extract_image_text(image)
            final_text = text if text.strip() else ocr_text
            final_source = source_type if text.strip() else "ocr"
            logger.info(
                f"[MULTIMODAL] Classified image as document_image (chars={len(final_text)}, method={final_source})"
            )
            return [
                MultimodalPage(
                    page_number=1,
                    text=final_text,
                    source_type=final_source,
                    input_type=InputType.DOCUMENT_IMAGE.value,
                    has_text=True,
                    has_visuals=True,
                )
            ]

        # Normal Photograph / Human Image / Scene / Diagram (No/low OCR text)
        # Rule: "No OCR text extracted" does NOT mean "no useful image information."
        logger.info("[MULTIMODAL] Classified image as visual_image / photograph. Initiating visual understanding.")

        # Send image directly to vision-capable model
        description, reason = ocr_vision_service.understand_image(image)

        if description and description.strip():
            logger.info(f"[MULTIMODAL] Vision model produced visual understanding ({len(description)} chars).")
            full_content = f"[Visual Content & Image Understanding]:\n{description}"
            return [
                MultimodalPage(
                    page_number=1,
                    text=full_content,
                    source_type="visual_understanding",
                    input_type=InputType.VISUAL_IMAGE.value,
                    has_text=False,
                    has_visuals=True,
                    visual_description=description,
                )
            ]

        # If vision model is unavailable or offline (e.g. Tesseract mode without API key),
        # preserve the image media entry gracefully rather than failing with "no content".
        fallback_descriptor = (
            f"[Visual Media: Image uploaded ({Path(path).name}). "
            f"No printed text was extracted by OCR. Visual understanding is available "
            f"when a vision model provider is configured.]"
        )
        logger.warning(
            f"[MULTIMODAL] Vision model unavailable (reason={reason}); registered visual descriptor."
        )
        return [
            MultimodalPage(
                page_number=1,
                text=fallback_descriptor,
                source_type="visual_media",
                input_type=InputType.VISUAL_IMAGE.value,
                has_text=False,
                has_visuals=True,
            )
        ]

    def process_docx(self, path: str) -> List[MultimodalPage]:
        import docx

        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        text = "\n".join(parts)
        return [
            MultimodalPage(
                page_number=1,
                text=text,
                source_type="text",
                input_type=InputType.STRUCTURED_DOC.value,
                has_text=True,
                has_visuals=False,
            )
        ]

    def process_txt(self, path: str) -> List[MultimodalPage]:
        with open(path, "r", errors="ignore") as f:
            text = f.read()
        return [
            MultimodalPage(
                page_number=1,
                text=text,
                source_type="text",
                input_type=InputType.STRUCTURED_DOC.value,
                has_text=True,
                has_visuals=False,
            )
        ]

    def process_csv(self, path: str) -> List[MultimodalPage]:
        import pandas as pd

        df = pd.read_csv(path)
        text = df.to_markdown(index=False)
        return [
            MultimodalPage(
                page_number=1,
                text=text,
                source_type="table",
                input_type=InputType.STRUCTURED_DOC.value,
                has_text=True,
                has_visuals=False,
            )
        ]

    def process_xlsx(self, path: str) -> List[MultimodalPage]:
        import pandas as pd

        sheets = pd.read_excel(path, sheet_name=None)
        pages = []
        for i, (name, df) in enumerate(sheets.items(), start=1):
            text = f"Sheet: {name}\n" + df.to_markdown(index=False)
            pages.append(
                MultimodalPage(
                    page_number=i,
                    text=text,
                    source_type="table",
                    input_type=InputType.STRUCTURED_DOC.value,
                    has_text=True,
                    has_visuals=False,
                )
            )
        return pages


multimodal_service = MultimodalService()
